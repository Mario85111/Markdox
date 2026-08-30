"""Testy jednostkowe torów konwersji (bez zależności od sieci/AI)."""
import io

import pymupdf
from docx import Document
from pptx import Presentation

from backend.app.converters import track_a
from backend.app.converters.detect import detect, get_ext
from backend.app.converters.service import convert_file


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Tytuł dokumentu", level=1)
    doc.add_paragraph("Zwykły akapit tekstu.")
    doc.add_heading("Sekcja", level=2)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pptx_bytes() -> bytes:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slajd tytułowy"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def test_get_ext():
    assert get_ext("plik.PDF") == "pdf"
    assert get_ext("obraz.JPG") == "jpg"
    assert get_ext("bez_rozszerzenia") == ""


def test_docx_to_markdown():
    md = track_a.docx_to_md(_docx_bytes())
    assert "# Tytuł dokumentu" in md
    assert "## Sekcja" in md
    assert "| A | B |" in md
    assert "| --- | --- |" in md


def test_pptx_title_becomes_heading():
    """Tytuł slajdu ma być nagłówkiem — to kotwica semantyczna chunka."""
    md = track_a.pptx_to_md(_pptx_bytes())
    assert "## Slajd tytułowy" in md
    assert "## Slajd 1" not in md            # etykieta bez treści
    assert md.count("Slajd tytułowy") == 1   # nie duplikujemy go jako punktora


def test_detect_track_a_for_docx():
    data = _docx_bytes()
    track, ext = detect(data, "test.docx")
    assert track == "A"
    assert ext == "docx"


def test_detect_image_is_track_b():
    track, ext = detect(b"\x89PNG\r\n", "skan.png")
    assert track == "B"


def test_text_passthrough_roundtrip():
    md = track_a.text_passthrough("# Cześć\n\nĄĘŚ".encode("utf-8"))
    assert "Cześć" in md and "ĄĘŚ" in md


def test_service_full_docx():
    result = convert_file(_docx_bytes(), "raport.docx", {})
    assert result["status"] == "ok"
    assert result["track"] == "A"
    assert result["markdown_filename"] == "raport.md"
    assert "# Tytuł dokumentu" in result["markdown"]


def test_service_unsupported_format():
    result = convert_file(b"dane", "plik.xyz", {})
    assert result["status"] == "error"
    assert "Nieobsługiwany" in result["error"]


# ---------- regresje jakości pod RAG ----------

def _two_column_pdf() -> bytes:
    """Realistyczny układ dwukolumnowy — akapity, nie pojedyncze etykiety."""
    doc = pymupdf.open()
    page = doc.new_page()
    y = 60
    for i in range(1, 4):
        left = f"LEWA-{i}: " + "tekst lewej kolumny o realistycznej długości. " * 3
        right = f"PRAWA-{i}: " + "tekst prawej kolumny o realistycznej długości. " * 3
        page.insert_textbox(pymupdf.Rect(50, y, 280, y + 90), left, fontsize=10)
        page.insert_textbox(pymupdf.Rect(310, y, 540, y + 90), right, fontsize=10)
        y += 100
    return doc.tobytes()


def test_pdf_reading_order_in_two_columns():
    """Kolumny nie mogą się przeplatać ani zlewać w jeden akapit.

    Sortowanie bloków po (y, x) dawało kolejność LEWA-1, PRAWA-1, LEWA-2…,
    czyli chunk łączył zdania z dwóch niezależnych kolumn.
    """
    md = track_a.pdf_text_to_md(_two_column_pdf())
    order = [md.index(f"{side}-{i}") for side in ("LEWA", "PRAWA") for i in (1, 2, 3)]
    assert order == sorted(order), "cała lewa kolumna musi poprzedzać prawą"
    for line in md.splitlines():
        assert not ("LEWA" in line and "PRAWA" in line), "zlepienie kolumn w jednej linii"


def test_pdf_headings_are_detected_from_font_size():
    """Nagłówki wymagają wyraźnej hierarchii wielkości względem tekstu."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((50, 60), "Umowa najmu lokalu", fontsize=24)
    page.insert_text((50, 110), "1. Przedmiot umowy", fontsize=16)
    y = 140
    for _ in range(6):
        page.insert_text((50, y), "Zwykły akapit treści umowy najmu lokalu.", fontsize=10)
        y += 18
    md = track_a.pdf_text_to_md(doc.tobytes())
    assert "# Umowa najmu lokalu" in md
    assert "## 1. Przedmiot umowy" in md


def test_pdf_tables_are_extracted_as_gfm():
    doc = pymupdf.open()
    page = doc.new_page()
    cells = [["Pozycja", "Kwota"], ["Czynsz", "2500"], ["Media", "400"]]
    for r, row in enumerate(cells):
        for c, val in enumerate(row):
            rect = pymupdf.Rect(50 + c * 150, 250 + r * 25, 50 + (c + 1) * 150, 250 + (r + 1) * 25)
            page.draw_rect(rect)
            page.insert_text((rect.x0 + 5, rect.y0 + 17), val, fontsize=10)
    md = track_a.pdf_text_to_md(doc.tobytes())
    assert "|Pozycja|Kwota|" in md.replace(" | ", "|")
    assert "Czynsz" in md and "2500" in md


def test_text_passthrough_decodes_iso8859_2_correctly():
    """Regresja: cp1250 nie zgłasza błędu, więc ISO-8859-2 dawało mojibake."""
    txt = "Zażółć gęślą jaźń"
    assert track_a.text_passthrough(txt.encode("iso-8859-2")) == txt
    assert track_a.text_passthrough(txt.encode("cp1250")) == txt
    assert track_a.text_passthrough(txt.encode("utf-16")) == txt


def test_text_passthrough_prefers_utf8():
    txt = "# Cześć\n\nĄĘŚ"
    assert track_a.text_passthrough(txt.encode("utf-8")) == txt
