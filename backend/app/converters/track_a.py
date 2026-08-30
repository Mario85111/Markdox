"""Tor A — dokumenty cyfrowe (docx, pptx, pdf-tekst, txt/md) → Markdown.

Deterministyczne parsery, bez AI, offline.
"""
import codecs
import io

import charset_normalizer
import pymupdf
import pymupdf4llm
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation


# ---------- helpery wspólne ----------

def _rows_to_gfm(rows: list[list[str]]) -> str:
    """Lista wierszy → tabela w składni GFM. Pierwszy wiersz = nagłówek."""
    rows = [[(c or "").strip().replace("\n", " ").replace("|", "\\|") for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |",
           "| " + " | ".join(["---"] * ncol) + " |"]
    for r in rows[1:]:
        out.append("| " + " | ".join(r) + " |")
    return "\n".join(out)


# ---------- DOCX ----------

def _iter_block_items(parent):
    """Iteruje akapity i tabele w kolejności występowania w dokumencie."""
    from docx.document import Document as _Doc
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P

    parent_elm = parent.element.body if isinstance(parent, _Doc) else parent
    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)


def _run_md(run) -> str:
    text = run.text or ""
    if not text.strip():
        return text
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()):]
    core = text.strip()
    if run.bold and run.italic:
        core = f"***{core}***"
    elif run.bold:
        core = f"**{core}**"
    elif run.italic:
        core = f"*{core}*"
    return f"{lead}{core}{trail}"


def _paragraph_md(p: Paragraph) -> str:
    text = "".join(_run_md(r) for r in p.runs).strip()
    if not text:
        return ""
    style = (p.style.name or "").lower() if p.style else ""
    if style.startswith("heading"):
        digits = "".join(ch for ch in style if ch.isdigit())
        level = min(max(int(digits) if digits else 1, 1), 6)
        return "#" * level + " " + text
    if style in ("title",):
        return "# " + text
    if style in ("subtitle",):
        return "## " + text
    if "list" in style:
        return "- " + text
    return text


def _docx_table_md(table: Table) -> str:
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    return _rows_to_gfm(rows)


def docx_to_md(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    parts: list[str] = []
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            md = _paragraph_md(block)
        else:
            md = _docx_table_md(block)
        if md:
            parts.append(md)
    return "\n\n".join(parts).strip()


# ---------- PPTX ----------

def _pptx_table_md(table) -> str:
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    return _rows_to_gfm(rows)


def pptx_to_md(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    slides_md: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        # Tytuł slajdu to kotwica semantyczna chunka — musi być nagłówkiem.
        # Wcześniej lądował wśród punktorów, a jedynym nagłówkiem było
        # "## Slajd N", czyli etykieta bez treści, powtarzalna w każdej talii.
        title_shape = slide.shapes.title
        title = (title_shape.text or "").strip() if title_shape is not None else ""
        # Porównujemy element XML, nie obiekt: `shapes.title` tworzy przy każdym
        # odwołaniu nowe opakowanie, więc `is` na kształcie zawsze dałoby False
        # i tytuł trafiłby do wyniku drugi raz, jako punktor.
        title_el = title_shape.element if title_shape is not None else None
        lines = [f"## {title}" if title else f"## Slajd {i}", ""]
        for shape in slide.shapes:
            if title_el is not None and shape.element is title_el:
                continue
            if shape.has_table:
                lines.append(_pptx_table_md(shape.table))
                lines.append("")
            elif shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if not text:
                        continue
                    level = para.level or 0
                    lines.append("  " * level + "- " + text)
                lines.append("")
        slides_md.append("\n".join(lines).strip())
    return "\n\n".join(s for s in slides_md if s).strip()


# ---------- PDF (warstwa tekstowa) ----------

def pdf_to_pages(file_bytes: bytes) -> list[str]:
    """PDF → Markdown, strona po stronie.

    Zamiast ręcznego sortowania bloków po (y, x) używamy pymupdf4llm:
    rozpoznaje kolejność czytania w układach wielokolumnowych, wyprowadza
    nagłówki z rozmiaru czcionki i wyciąga tabele jako GFM. Poprzednie
    sortowanie przeplatało kolumny i sklejało niepowiązane zdania
    w jeden akapit, co w RAG dawało chunki bez sensu.
    """
    with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
        chunks = pymupdf4llm.to_markdown(
            doc,
            page_chunks=True,     # potrzebne do numerów stron w prowenancji
            ignore_images=True,   # placeholdery obrazów to szum dla RAG
            show_progress=False,
        )
    return [(c.get("text") or "").strip() for c in chunks]


def pdf_text_to_md(file_bytes: bytes) -> str:
    """Wersja płaska — strony rozdzielone poziomą linią."""
    return "\n\n---\n\n".join(p for p in pdf_to_pages(file_bytes) if p).strip()


# ---------- TXT / MD ----------

_BOMS = (
    (codecs.BOM_UTF8, "utf-8-sig"),
    (codecs.BOM_UTF32_LE, "utf-32"), (codecs.BOM_UTF32_BE, "utf-32"),
    (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16"),
)

# Kandydaci realistyczni dla dokumentów polskich i angielskich.
_CANDIDATES = ("cp1250", "iso-8859-2", "cp1252", "cp852")

_PL_DIACRITICS = set("ąćęłńóśźżĄĆĘŁŃÓŚŹŻ")
# Znaki dopuszczalne w prozie poza literami ASCII, cyframi i białymi znakami.
_PUNCT_OK = set(".,;:!?-–—()[]{}\"'„”“…/\\@#%&*+=<>|~^$§°€£©®№ ")


def _plausibility(text: str) -> int:
    """Ocena, na ile tekst wygląda na polską lub angielską prozę.

    Detektory ogólnego przeznaczenia mylą kodowania jednobajtowe:
    dla polskiego zdania wskazują cp1257 czy iso8859-10, bo statystyka
    bajtów jest niemal identyczna. Wiedza o docelowym języku rozstrzyga
    to, czego statystyka rozstrzygnąć nie potrafi.
    """
    score = 0
    for ch in text:
        if ch in _PL_DIACRITICS:
            score += 3                      # mocny sygnał poprawnego dekodowania
        elif ch.isascii() and (ch.isalnum() or ch.isspace()) or ch in _PUNCT_OK:
            continue
        elif ch.isprintable():
            score -= 4                      # znak spoza repertuaru — pewnie mojibake
        else:
            score -= 10                     # znak sterujący w pliku tekstowym
    return score


def text_passthrough(file_bytes: bytes) -> str:
    """Dekoduje plik tekstowy — bez cichego psucia znaków.

    Poprzednia wersja próbowała kodowań po kolei, ale cp1250 jest
    jednobajtowe i praktycznie nigdy nie zgłasza błędu, więc plik
    w ISO-8859-2 stawał się mojibake bez ostrzeżenia.

    Kolejność: BOM → UTF-8 → wybór kandydata po ocenie wiarygodności
    → detektor ogólny jako ostatnia deska ratunku → głośny błąd.
    """
    for bom, enc in _BOMS:
        if file_bytes.startswith(bom):
            return file_bytes.decode(enc).strip()
    try:
        return file_bytes.decode("utf-8").strip()
    except UnicodeDecodeError:
        pass

    scored = []
    for enc in _CANDIDATES:
        try:
            text = file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
        scored.append((_plausibility(text), text))
    if scored:
        best_score, best_text = max(scored, key=lambda pair: pair[0])
        if best_score >= 0:
            return best_text.strip()

    fallback = charset_normalizer.from_bytes(file_bytes).best()
    if fallback is None:
        raise ValueError(
            "Nie rozpoznano kodowania pliku tekstowego — zapisz go w UTF-8."
        )
    return str(fallback).strip()


# ---------- dispatcher ----------

def convert_track_a(file_bytes: bytes, ext: str) -> str:
    if ext == "docx":
        return docx_to_md(file_bytes)
    if ext == "pptx":
        return pptx_to_md(file_bytes)
    if ext == "pdf":
        return pdf_text_to_md(file_bytes)
    if ext in ("txt", "md", "markdown"):
        return text_passthrough(file_bytes)
    raise ValueError(f"Tor A nie obsługuje rozszerzenia .{ext}")
